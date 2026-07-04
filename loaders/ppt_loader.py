import logging
from pathlib import Path 
from pptx import Presentation 
from langchain_core.documents import Document 
from typing import List
logger = logging.getLogger(__name__)

def load_ppt(file_path : str, doc_id : str ) -> List[Document]:
    path = Path(file_path)
    if not path.exists() :
        raise FileNotFoundError(f'presentation file is not found at {path.name}')
    if path.suffix.lower == '.ppt':
        raise ValueError(
            f"{path.name} is a legacy .ppt(binary) file. Only modern .pptx "
            f"format is supported - please convert and reupload it "
        )
    try : 
        prs = Presentation(file_path)
    except Exception as e : 
        logger.error(f"Failed to open presentattion file at {path.name} : {e}")
        raise ValueError(f"{path.name} is not a valid .pptx file or is corrupted")
    docs :List[Document] =  []
    total_slides = len(prs.slides)
    for i  , slide in enumerate(prs.slides):
        slide_Number = i + 1
        title_text = ""
        body_text = ""
        table_text = []
        for shape in slide.shapes: 
            if shape.has_text_frame:
                text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if not text : 
                continue
            if shape  == slide.shapes.title:
                title_text = text 
            else:
                body_text.append(text)
            if shape.has_table : 
                table = shape.table
                rows_text = []
                for row in table.rows : 
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        rows_text.append("|".join(cells))
                if rows_text:
                    table_text.append("\n".join(rows_text))
        notes_text = ""
        if slide.has_notes_slide: 
            notes_text = (slide.notes_slide.notes_text_frame.text or "").strip()
        content_parts =  []
        if title_text:
            content_parts.append(f'Slide Title : {title_text}')
        if body_text :
            content_parts.append("\n".join(body_text))
        if table_text:
            content_parts.append("[TABLE]\n" + "\n\n".join(table_text) + "\n[TABLE]")
        if notes_text:
            content_parts.append(f"[SPEAKER NOTES]\n{notes_text}\n[/SPEAKER NOTES]")
        fullcontent  =  "\n\n".join(content_parts.strip())
        if not fullcontent: 
            logger.debug(f'Slide Number {slide_Number} is Empty , skipping')
            continue
        docs.append(
            Document(
                page_content=fullcontent,
                metadata={
                    "source": path.name,
                    "doc_id": doc_id,
                    "slide_number": slide_Number,
                    "total_slides": total_slides,
                    "slide_title": title_text or None,
                    "has_speaker_notes": bool(notes_text),
                    "file_type": "pptx",
                },
            ))


    if not docs :
        raise ValueError(f"PPTX {path.name} has no extractable text on any of its {total_slides} slides")
    logger.info(f"Loaded PPTX '{path.name}': {len(docs)}/{total_slides} non-empty slides")
    return docs 
    
             
        
    
        